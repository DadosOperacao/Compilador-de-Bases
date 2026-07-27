import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError
from .utils import sanitizar_valor_excel

def read_csv_semicolon(path):
    return pd.read_csv(path, sep=";", low_memory=False, dtype=str)

def read_excel_default(path):
    return pd.read_excel(path)

def save_excel_safe(df, output_file, sheet_name="Base Consolidada"):
    output_file.parent.mkdir(parents=True, exist_ok=True)
    if output_file.exists():
        output_file.unlink()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for c, name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c, value=sanitizar_valor_excel(name))
    for r, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c, value in enumerate(row, start=1):
            try:
                ws.cell(row=r, column=c, value=sanitizar_valor_excel(value))
            except IllegalCharacterError:
                ws.cell(row=r, column=c, value="")
    wb.save(output_file)
